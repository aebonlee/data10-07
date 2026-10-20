"""기획서 Markdown → Word(docx). 실행: python3 scripts/md2docx.py docs/01_프로젝트_기획서.md
python-docx 가 필요합니다(pip install python-docx). 제목·표·글머리·번호·굵게·코드 블록(mermaid 는 안내 문구로)만 다룹니다."""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = '맑은 고딕'


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(par, text, size=10.5):
    # **굵게**, `코드` 만 처리
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            set_font(par.add_run(part[2:-2]), size, True)
        elif part.startswith('`') and part.endswith('`'):
            set_font(par.add_run(part[1:-1]), size, None, '174A84')
        else:
            set_font(par.add_run(part), size)


def cells(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def convert(src: Path, dst: Path):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.font.size = Pt(10.5)
    lines = src.read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith('```'):
            lang = s[3:].strip()
            body = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                body.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            if lang == 'mermaid':
                add_inline(p, '(흐름도 — Markdown 판(docs/01_프로젝트_기획서.md)에서 그림으로 볼 수 있습니다)', 9.5)
            else:
                set_font(p.add_run('\n'.join(body)), 9)
            continue
        if s.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                if not re.match(r'^\|?\s*:?-{2,}', lines[i].strip()):
                    rows.append(cells(lines[i]))
                i += 1
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.style = 'Table Grid'
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, row in enumerate(rows):
                for c in range(ncol):
                    cell = t.cell(r, c)
                    cell.text = ''
                    add_inline(cell.paragraphs[0], row[c] if c < len(row) else '', 9.5)
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
            continue
        m = re.match(r'^(#{1,4})\s+(.*)$', s)
        if m:
            level = len(m.group(1))
            hp = doc.add_heading(level=min(level, 3))
            set_font(hp.add_run(m.group(2)), {1: 16, 2: 13.5, 3: 12, 4: 11}[level], True, '13274B')
        elif s == '---':
            pass
        elif re.match(r'^[-*]\s+', s):
            add_inline(doc.add_paragraph(style='List Bullet'), re.sub(r'^[-*]\s+', '', s))
        elif re.match(r'^\d+\.\s+', s):
            # 번호는 글자 그대로 둡니다 — Word 의 자동 번호는 문서 전체에서 이어져 10장 목록이 6부터 시작하는 일이 생깁니다
            np_ = doc.add_paragraph()
            np_.paragraph_format.left_indent = Pt(18)
            np_.paragraph_format.first_line_indent = Pt(-14)
            add_inline(np_, s)
        elif s:
            add_inline(doc.add_paragraph(), s)
        i += 1
    doc.save(dst)


if __name__ == '__main__':
    src = Path(sys.argv[1] if len(sys.argv) > 1 else 'docs/01_프로젝트_기획서.md')
    dst = src.with_suffix('.docx')
    convert(src, dst)
    print(dst)
